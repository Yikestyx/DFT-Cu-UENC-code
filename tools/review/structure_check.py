import os
import subprocess
from collections import OrderedDict
import tkinter as tk
from tkinter import messagebox, simpledialog, ttk

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageTk

# ==================== Config ====================
EXCEL_PATH = r"D:\lyf\calculation\Cu-old\couple\Cu\energy_clean.xlsx"
SHEET_NAME = None                      # None -> first sheet
NAME_COLUMN = "产物"                    # fallback to first column if missing
COMMENT_COLUMN = "_Comment"            # comment column (before MARK_COLUMN)
MARK_COLUMN = "_RatingMark"            # hidden rating column

BASE_DIR = r"D:\lyf\calculation\Cu-old\couple\Cu"
IMAGE_FOLDER_CANDIDATES = ["img", "image", "images"]
CONTCAR_DIR = os.path.join(BASE_DIR, "CONTCAR")
VESTA_PATH = r"C:\Users\logic\Desktop\VESTA-win64\VESTA.exe"

COLOR_GOOD = "c6efce"
COLOR_BAD = "ffc7ce"
WINDOW_TOPMOST = False                 # user requested: disable topmost

DISPLAY_SIZE = (460, 360)              # prefetch resize target
CACHE_MAX_MB = 180                     # global cache memory cap
PREFETCH_RADIUS = 2                    # preload neighbors
MAX_VIEWED_STRUCTURES = 5              # keep cache only for recent viewed structures
# ================================================


class LruPhotoCache:
    def __init__(self, max_bytes: int):
        self.max_bytes = max_bytes
        self.cache = OrderedDict()  # key -> (photo, bytes)
        self.total_bytes = 0

    @staticmethod
    def _estimate_bytes(photo: ImageTk.PhotoImage) -> int:
        return int(photo.width() * photo.height() * 4)

    def get(self, key: str):
        if key not in self.cache:
            return None
        photo, size = self.cache.pop(key)
        self.cache[key] = (photo, size)
        return photo

    def put(self, key: str, photo: ImageTk.PhotoImage):
        if key in self.cache:
            _, old = self.cache.pop(key)
            self.total_bytes -= old
        size = self._estimate_bytes(photo)
        self.cache[key] = (photo, size)
        self.total_bytes += size
        self._evict()

    def _evict(self):
        while self.total_bytes > self.max_bytes and self.cache:
            _, (_, size) = self.cache.popitem(last=False)
            self.total_bytes -= size

    def remove_where(self, predicate):
        keys = [k for k in self.cache.keys() if predicate(k)]
        for key in keys:
            _, size = self.cache.pop(key)
            self.total_bytes -= size


class StructureImageMarkerApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("结构筛选工具（图像快速判定）")
        self.root.geometry("1540x880")
        if WINDOW_TOPMOST:
            self.root.attributes("-topmost", True)

        self.excel_path = EXCEL_PATH
        self.sheet_name = SHEET_NAME

        self.wb = None
        self.ws = None
        self.name_col_idx = None
        self.comment_col_idx = None
        self.mark_col_idx = None

        self.all_rows = []  # [(row_idx, struct_name)]
        self.current_idx = -1

        self.image_dir = ""
        self.image_map = {}  # struct_name -> {front/side/top: path}
        self.photo_cache = LruPhotoCache(max_bytes=CACHE_MAX_MB * 1024 * 1024)
        self.current_views = {"front": "", "side": "", "top": ""}
        self.viewed_structures = []

        self.color_good = PatternFill(start_color=COLOR_GOOD, end_color=COLOR_GOOD, fill_type="solid")
        self.color_bad = PatternFill(start_color=COLOR_BAD, end_color=COLOR_BAD, fill_type="solid")

        self._build_ui()
        self._bind_shortcuts()

        if os.path.exists(self.excel_path):
            self._load_all(self.excel_path)
        else:
            self.status_var.set(f"默认Excel不存在：{self.excel_path}")

    def _build_ui(self):
        top = tk.Frame(self.root)
        top.pack(fill=tk.X, padx=8, pady=6)

        tk.Button(top, text="重新加载", command=self.reload_data).pack(side=tk.LEFT, padx=4)
        tk.Button(top, text="生成Selected表", command=self.generate_selected_sheet, bg="#d9edf7").pack(side=tk.LEFT, padx=4)

        self.image_dir_var = tk.StringVar(value="图像目录: 未检测")
        tk.Label(top, textvariable=self.image_dir_var, anchor="w").pack(side=tk.LEFT, padx=14)

        info = tk.LabelFrame(self.root, text="当前结构", padx=8, pady=8)
        info.pack(fill=tk.X, padx=8)

        tk.Label(info, text="结构名称:").grid(row=0, column=0, sticky="w", pady=3)
        self.name_var = tk.StringVar()
        tk.Entry(info, textvariable=self.name_var, state="readonly", width=50).grid(row=0, column=1, sticky="w", padx=5)

        tk.Label(info, text="当前评价:").grid(row=1, column=0, sticky="w", pady=3)
        self.rating_var = tk.StringVar()
        tk.Entry(info, textvariable=self.rating_var, state="readonly", width=50).grid(row=1, column=1, sticky="w", padx=5)

        tk.Label(info, text="当前批注:").grid(row=2, column=0, sticky="w", pady=3)
        self.comment_var = tk.StringVar()
        tk.Entry(info, textvariable=self.comment_var, state="readonly", width=50).grid(row=2, column=1, sticky="w", padx=5)

        tk.Label(info, text="图像匹配:").grid(row=3, column=0, sticky="w", pady=3)
        self.match_var = tk.StringVar()
        tk.Entry(info, textvariable=self.match_var, state="readonly", width=80).grid(row=3, column=1, columnspan=4, sticky="we", padx=5)

        nav = tk.Frame(info)
        nav.grid(row=0, column=2, rowspan=2, padx=10)
        tk.Button(nav, text="上一条", command=self.go_prev, width=10).pack(side=tk.LEFT, padx=4)
        tk.Button(nav, text="下一条", command=self.go_next, width=10).pack(side=tk.LEFT, padx=4)
        tk.Button(nav, text="VESTA", command=self.open_current_in_vesta, width=10).pack(side=tk.LEFT, padx=4)

        rate = tk.Frame(info)
        rate.grid(row=0, column=3, rowspan=2, padx=10)
        tk.Button(rate, text="Good", command=lambda: self.set_rating("Good"), bg=f"#{COLOR_GOOD}", width=10).pack(side=tk.LEFT, padx=4)
        tk.Button(rate, text="Bad", command=lambda: self.set_rating("Bad"), bg=f"#{COLOR_BAD}", width=10).pack(side=tk.LEFT, padx=4)
        tk.Button(rate, text="跳过", command=self.go_next, width=10).pack(side=tk.LEFT, padx=4)

        comment_btn = tk.Frame(info)
        comment_btn.grid(row=2, column=2, columnspan=2, padx=10, sticky="w")
        tk.Button(comment_btn, text="未成键", command=lambda: self.set_comment("未成键"), width=10).pack(side=tk.LEFT, padx=4)
        tk.Button(comment_btn, text="解离", command=lambda: self.set_comment("解离"), width=10).pack(side=tk.LEFT, padx=4)
        tk.Button(comment_btn, text="自定义批注", command=self.ask_custom_comment, width=12).pack(side=tk.LEFT, padx=4)
        tk.Button(comment_btn, text="删除批注", command=self.clear_comment, width=10).pack(side=tk.LEFT, padx=4)

        images = tk.LabelFrame(self.root, text="结构图（三视图自动切换）", padx=8, pady=8)
        images.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.image_canvases = {}
        for i, view in enumerate(["front", "side", "top"]):
            col = tk.Frame(images)
            col.grid(row=0, column=i, sticky="nsew", padx=6)
            images.columnconfigure(i, weight=1)

            tk.Label(col, text=view.upper()).pack(anchor="center", pady=(0, 6))
            canvas = tk.Canvas(col, bg="#f5f5f5", highlightthickness=1, highlightbackground="#cfcfcf")
            canvas.pack(fill=tk.BOTH, expand=True)
            canvas.bind("<Configure>", lambda _, v=view: self._rerender_view(v))
            self.image_canvases[view] = canvas

        status = tk.Frame(self.root)
        status.pack(fill=tk.X, padx=8, pady=(0, 8))

        self.progress = ttk.Progressbar(status, orient=tk.HORIZONTAL, length=420, mode="determinate")
        self.progress.pack(side=tk.LEFT, padx=4)

        self.status_var = tk.StringVar(value="就绪")
        tk.Label(status, textvariable=self.status_var).pack(side=tk.LEFT, padx=8)

        self.cache_var = tk.StringVar(value="缓存: 0.0 MB")
        tk.Label(status, textvariable=self.cache_var).pack(side=tk.RIGHT, padx=8)

        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)

    def _bind_shortcuts(self):
        self.root.bind("<Left>", lambda _: self.go_prev())
        self.root.bind("<Right>", lambda _: self.go_next())
        self.root.bind("<Up>", lambda _: self.set_rating("Good"))
        self.root.bind("<Down>", lambda _: self.set_rating("Bad"))

    def reload_data(self):
        if self.wb:
            self.wb.close()
        self._load_all(self.excel_path)

    def _resolve_image_dir(self, excel_path: str) -> str:
        base = os.path.dirname(os.path.abspath(excel_path))
        for name in IMAGE_FOLDER_CANDIDATES:
            cand = os.path.join(base, name)
            if os.path.isdir(cand):
                return cand
        for name in IMAGE_FOLDER_CANDIDATES:
            cand = os.path.join(BASE_DIR, name)
            if os.path.isdir(cand):
                return cand
        return ""

    def _load_all(self, excel_path: str):
        try:
            self.wb = load_workbook(excel_path)
            if self.sheet_name and self.sheet_name in self.wb.sheetnames:
                self.ws = self.wb[self.sheet_name]
            else:
                self.ws = self.wb[self.wb.sheetnames[0]]

            self._resolve_columns()
            self._load_rows()

            self.image_dir = self._resolve_image_dir(excel_path)
            self.image_map = self._build_image_map(self.image_dir) if self.image_dir else {}
            self.image_dir_var.set(f"图像目录: {self.image_dir or '未找到（请检查 img/image 目录）'}")

            self.current_idx = 0 if self.all_rows else -1
            self.viewed_structures.clear()
            self.update_display(auto_prefetch=True)
            self.status_var.set(f"已加载 {len(self.all_rows)} 个结构（工作表: {self.ws.title}）")
        except Exception as e:
            messagebox.showerror("错误", f"读取Excel失败:\n{e}")

    def _resolve_columns(self):
        self.name_col_idx = None
        self.comment_col_idx = None
        self.mark_col_idx = None

        for col in range(1, self.ws.max_column + 1):
            header = self.ws.cell(row=1, column=col).value
            if header == NAME_COLUMN:
                self.name_col_idx = col
            elif header == COMMENT_COLUMN:
                self.comment_col_idx = col
            elif header == MARK_COLUMN:
                self.mark_col_idx = col

        if self.name_col_idx is None:
            self.name_col_idx = 1

        if self.mark_col_idx is None:
            self.mark_col_idx = self.ws.max_column + 1
            self.ws.cell(row=1, column=self.mark_col_idx, value=MARK_COLUMN)
        self.ws.column_dimensions[get_column_letter(self.mark_col_idx)].hidden = True

        if self.comment_col_idx is None:
            self.ws.insert_cols(self.mark_col_idx, amount=1)
            self.comment_col_idx = self.mark_col_idx
            self.ws.cell(row=1, column=self.comment_col_idx, value=COMMENT_COLUMN)
            self.mark_col_idx += 1
        elif self.comment_col_idx > self.mark_col_idx:
            self.ws.insert_cols(self.mark_col_idx, amount=1)
            self.ws.cell(row=1, column=self.mark_col_idx, value=COMMENT_COLUMN)
            old_comment_idx = self.comment_col_idx + 1
            for r in range(2, self.ws.max_row + 1):
                self.ws.cell(row=r, column=self.mark_col_idx, value=self.ws.cell(row=r, column=old_comment_idx).value)
            self.ws.delete_cols(old_comment_idx, amount=1)
            self.comment_col_idx = self.mark_col_idx
            self.mark_col_idx = self.comment_col_idx + 1

        # user requested: comment column should NOT be hidden
        self.ws.column_dimensions[get_column_letter(self.comment_col_idx)].hidden = False

    def _load_rows(self):
        self.all_rows = []
        for r in range(2, self.ws.max_row + 1):
            v = self.ws.cell(row=r, column=self.name_col_idx).value
            if v is None:
                continue
            name = str(v).strip()
            if name:
                self.all_rows.append((r, name))

    @staticmethod
    def _split_view_from_filename(filename: str):
        lower = filename.lower()
        for view in ("front", "side", "top"):
            token = f"_{view}.png"
            if lower.endswith(token):
                struct = filename[:-len(token)]
                return struct, view
        return None, None

    def _build_image_map(self, image_dir: str):
        mapping = {}
        if not image_dir:
            return mapping
        for fn in os.listdir(image_dir):
            if not fn.lower().endswith(".png"):
                continue
            struct, view = self._split_view_from_filename(fn)
            if not struct:
                continue
            mapping.setdefault(struct, {})[view] = os.path.join(image_dir, fn)
        return mapping

    @staticmethod
    def _clean_struct_name(name: str) -> str:
        return name[1:] if name.startswith("*") else name

    def _load_photo(self, path: str, max_size):
        max_w, max_h = max_size
        if max_w <= 1 or max_h <= 1:
            max_w, max_h = DISPLAY_SIZE
        key = f"{path}|{max_w}x{max_h}"
        cached = self.photo_cache.get(key)
        if cached is not None:
            return cached
        img = Image.open(path).convert("RGB")
        img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
        photo = ImageTk.PhotoImage(img)
        self.photo_cache.put(key, photo)
        return photo

    def _set_canvas_image(self, canvas: tk.Canvas, path: str):
        canvas.delete("all")
        if not path or not os.path.exists(path):
            w = max(canvas.winfo_width(), 160)
            h = max(canvas.winfo_height(), 120)
            canvas.create_text(w // 2, h // 2, text="无图", fill="#666666")
            canvas.image = None
            return
        try:
            w = max(canvas.winfo_width() - 8, 50)
            h = max(canvas.winfo_height() - 8, 50)
            photo = self._load_photo(path, (w, h))
            canvas.create_image(canvas.winfo_width() // 2, canvas.winfo_height() // 2, image=photo, anchor="center")
            canvas.image = photo
        except Exception:
            w = max(canvas.winfo_width(), 160)
            h = max(canvas.winfo_height(), 120)
            canvas.create_text(w // 2, h // 2, text="读图失败", fill="#cc0000")
            canvas.image = None

    def _rerender_view(self, view: str):
        canvas = self.image_canvases.get(view)
        if canvas is None:
            return
        self._set_canvas_image(canvas, self.current_views.get(view, ""))

    def _update_cache_status(self):
        used_mb = self.photo_cache.total_bytes / (1024 * 1024)
        self.cache_var.set(f"缓存: {used_mb:.1f} MB")

    def _remember_viewed_structure(self, struct_name: str):
        if not struct_name:
            return
        if struct_name in self.viewed_structures:
            self.viewed_structures.remove(struct_name)
        self.viewed_structures.append(struct_name)
        while len(self.viewed_structures) > MAX_VIEWED_STRUCTURES:
            removed = self.viewed_structures.pop(0)
            views = self.image_map.get(removed, {})
            prefixes = [f"{p}|" for p in views.values() if p]
            if prefixes:
                # noinspection PyUnresolvedReferences
                self.photo_cache.remove_where(
                    lambda key, prefixes=prefixes: any(key.startswith(prefix) for prefix in prefixes)
                )
        self._update_cache_status()

    def _prefetch_neighbors(self):
        if not self.all_rows:
            return
        indices = []
        for step in range(1, PREFETCH_RADIUS + 1):
            prev_i = self.current_idx - step
            next_i = self.current_idx + step
            if 0 <= prev_i < len(self.all_rows):
                indices.append(prev_i)
            if 0 <= next_i < len(self.all_rows):
                indices.append(next_i)
        for idx in indices:
            _, raw = self.all_rows[idx]
            name = self._clean_struct_name(raw)
            views = self.image_map.get(name, {})
            for view in ("front", "side", "top"):
                p = views.get(view)
                if p and os.path.exists(p):
                    try:
                        self._load_photo(p, DISPLAY_SIZE)
                    except Exception:
                        pass
        self._update_cache_status()

    def update_display(self, auto_prefetch=False):
        if self.current_idx < 0 or self.current_idx >= len(self.all_rows):
            self.name_var.set("")
            self.rating_var.set("")
            self.comment_var.set("")
            self.match_var.set("")
            for view in ("front", "side", "top"):
                self.current_views[view] = ""
                self._set_canvas_image(self.image_canvases[view], "")
            self.progress["value"] = 100 if self.all_rows else 0
            self.status_var.set("无可显示结构" if not self.all_rows else "已到末尾")
            return

        row, raw_name = self.all_rows[self.current_idx]
        name = self._clean_struct_name(raw_name)
        self._remember_viewed_structure(name)
        self.name_var.set(name)

        mark = self.ws.cell(row=row, column=self.mark_col_idx).value
        self.rating_var.set(mark if mark else "未评价")

        comment = self.ws.cell(row=row, column=self.comment_col_idx).value
        self.comment_var.set(comment if comment else "")

        views = self.image_map.get(name, {})
        for view in ("front", "side", "top"):
            self.current_views[view] = views.get(view, "")
            self._set_canvas_image(self.image_canvases[view], self.current_views[view])

        matched = [v for v in ("front", "side", "top") if views.get(v)]
        if matched:
            self.match_var.set(f"已匹配: {', '.join(matched)}")
        else:
            self.match_var.set("未匹配到图片（检查文件名是否为 结构名_front/side/top.png）")

        self.progress["value"] = ((self.current_idx + 1) / len(self.all_rows)) * 100
        self.status_var.set(f"进度: {self.current_idx + 1}/{len(self.all_rows)}")

        if auto_prefetch:
            self.root.after_idle(self._prefetch_neighbors)
            self.root.after(40, lambda: [self._rerender_view(v) for v in ("front", "side", "top")])

    def set_rating(self, rating: str):
        if self.current_idx < 0 or self.current_idx >= len(self.all_rows):
            return
        row, _ = self.all_rows[self.current_idx]
        self.ws.cell(row=row, column=self.mark_col_idx).value = rating
        fill = self.color_good if rating == "Good" else self.color_bad
        for col in range(1, self.ws.max_column + 1):
            self.ws.cell(row=row, column=col).fill = fill
        try:
            self.wb.save(self.excel_path)
        except Exception as e:
            messagebox.showerror("错误", f"保存Excel失败:\n{e}")
            return
        self.go_next()

    def set_comment(self, comment: str):
        if self.current_idx < 0 or self.current_idx >= len(self.all_rows):
            return
        row, _ = self.all_rows[self.current_idx]
        self.ws.cell(row=row, column=self.comment_col_idx).value = comment
        try:
            self.wb.save(self.excel_path)
            self.comment_var.set(comment)
        except Exception as e:
            messagebox.showerror("错误", f"保存批注失败:\n{e}")

    def ask_custom_comment(self):
        value = simpledialog.askstring("自定义批注", "请输入批注内容：", parent=self.root)
        if value is None:
            return
        value = value.strip()
        if not value:
            return
        self.set_comment(value)

    def clear_comment(self):
        self.set_comment("")

    def go_prev(self):
        if self.current_idx > 0:
            self.current_idx -= 1
            self.update_display(auto_prefetch=True)

    def go_next(self):
        if self.current_idx < len(self.all_rows) - 1:
            self.current_idx += 1
            self.update_display(auto_prefetch=True)
        else:
            self.update_display(auto_prefetch=False)
            messagebox.showinfo("提示", "已经是最后一个结构")

    def open_current_in_vesta(self):
        if self.current_idx < 0 or self.current_idx >= len(self.all_rows):
            messagebox.showwarning("提示", "当前没有可打开的结构")
            return

        if not os.path.exists(VESTA_PATH):
            messagebox.showerror("错误", f"未找到VESTA可执行文件:\n{VESTA_PATH}")
            return

        _, raw_name = self.all_rows[self.current_idx]
        name = self._clean_struct_name(raw_name)
        cif_path = os.path.join(CONTCAR_DIR, f"{name}.cif")

        if not os.path.exists(cif_path):
            messagebox.showerror("错误", f"未找到当前结构CIF文件:\n{cif_path}")
            return

        try:
            subprocess.Popen([VESTA_PATH, cif_path], cwd=os.path.dirname(VESTA_PATH))
        except Exception as e:
            messagebox.showerror("错误", f"调用VESTA失败:\n{e}")

    def generate_selected_sheet(self):
        if self.ws is None:
            messagebox.showerror("错误", "请先加载Excel")
            return
        good_rows = []
        for row in range(2, self.ws.max_row + 1):
            val = self.ws.cell(row=row, column=self.mark_col_idx).value
            if str(val).strip() == "Good":
                good_rows.append(row)
        if not good_rows:
            messagebox.showwarning("提示", "没有标记为 Good 的结构")
            return
        if "Selected" in self.wb.sheetnames:
            old = self.wb["Selected"]
            self.wb.remove(old)
        new_ws = self.wb.create_sheet("Selected")
        cols = [c for c in range(1, self.ws.max_column + 1) if c != self.mark_col_idx]
        for nc, oc in enumerate(cols, start=1):
            new_ws.cell(row=1, column=nc, value=self.ws.cell(row=1, column=oc).value)
        for nr, orow in enumerate(good_rows, start=2):
            for nc, oc in enumerate(cols, start=1):
                new_ws.cell(row=nr, column=nc, value=self.ws.cell(row=orow, column=oc).value)
        try:
            self.wb.save(self.excel_path)
            messagebox.showinfo("成功", f"已生成 Selected 工作表，共 {len(good_rows)} 条")
        except Exception as e:
            messagebox.showerror("错误", f"保存失败:\n{e}")

    def on_closing(self):
        if self.wb:
            self.wb.close()
        self.root.destroy()


if __name__ == "__main__":
    root = tk.Tk()
    app = StructureImageMarkerApp(root)
    root.mainloop()
