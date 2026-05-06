import os
import re
from typing import Dict, List, Optional, Tuple

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


class StructurePPTGenerator:
    """
    从目录读取结构图像并生成 PPT。

    支持两种输入：
    1) 单结构模式: {prefix}_top/front/side.png
    2) 对比模式: {prefix}_POSCAR_top/front/side.png + {prefix}_CONTCAR_top/front/side.png

    当前排版：
    - save_both_views=False: 每个结构显示 top + (front 或 side 之一)
    - save_both_views=True:  每个结构显示 top + front + side（front 和 side 水平并排）
    """

    # ==================== 配置区（尺寸/位置控制，单位：英寸或比例） ====================
    SLIDE_WIDTH = 13.33
    SLIDE_HEIGHT = 7.5
    OUTER_PADDING = 0.15
    COLUMN_GAP = 0.18
    TITLE_HEIGHT_RATIO = 0.08
    TITLE_TOP_OFFSET_RATIO = 0.00
    TITLE_LEFT_OFFSET_RATIO = 0.00
    TITLE_WIDTH_RATIO = 1.00
    TITLE_FONT_SIZE = 18
    CONTENT_TOP_GAP_RATIO = 0.02
    PANEL_GAP_RATIO = 0.06
    LABEL_HEIGHT_RATIO = 0.06
    PANEL_LABEL_FONT_SIZE = 12
    IMAGE_GAP_RATIO = 0.04
    IMAGE_WIDTH_RATIO = 1.00
    TOP_IMAGE_HEIGHT_RATIO = 0.37
    BOTTOM_IMAGE_HEIGHT_RATIO = 0.43
    HORIZONTAL_VIEW_GAP_RATIO = 0.04
    SHOW_COLUMN_DIVIDER = False
    DIVIDER_X_OFFSET = 0.0
    DIVIDER_Y_OFFSET_RATIO = 0.00
    DIVIDER_HEIGHT_RATIO = 1.00
    DIVIDER_WIDTH_PT = 1.0
    DIVIDER_COLOR_RGB = (160, 160, 160)
    # =============================================================================

    def __init__(self, image_dir: str, save_both_views: bool = False):
        self.image_dir = os.path.abspath(image_dir)
        self.save_both_views = save_both_views
        self.structures: List[Dict] = []
        self._scan_images()

    def _scan_images(self):
        """扫描目录并按结构前缀聚合图像。"""
        if not os.path.isdir(self.image_dir):
            raise ValueError(f"目录不存在: {self.image_dir}")

        all_files = [f for f in os.listdir(self.image_dir) if f.lower().endswith(".png")]
        pattern = re.compile(r"^(.*?)_(top|front|side)\.png$", re.IGNORECASE)

        groups: Dict[str, Dict] = {}

        for filename in all_files:
            m = pattern.match(filename)
            if not m:
                continue

            prefix_full = m.group(1)
            view_type = m.group(2).lower()
            full_path = os.path.join(self.image_dir, filename)

            label = None
            base_prefix = prefix_full
            if prefix_full.endswith("_POSCAR"):
                label = "POSCAR"
                base_prefix = prefix_full[: -len("_POSCAR")]
            elif prefix_full.endswith("_CONTCAR"):
                label = "CONTCAR"
                base_prefix = prefix_full[: -len("_CONTCAR")]

            if base_prefix not in groups:
                groups[base_prefix] = {
                    "prefix": base_prefix,
                    "single": {"top": None, "front": None, "side": None},
                    "pair": {
                        "POSCAR": {"top": None, "front": None, "side": None},
                        "CONTCAR": {"top": None, "front": None, "side": None},
                    },
                }

            if label is None:
                groups[base_prefix]["single"][view_type] = full_path
            else:
                groups[base_prefix]["pair"][label][view_type] = full_path

        for prefix, data in groups.items():
            pair = data["pair"]
            pos_ok = pair["POSCAR"]["top"] is not None and (
                pair["POSCAR"]["front"] is not None or pair["POSCAR"]["side"] is not None
            )
            con_ok = pair["CONTCAR"]["top"] is not None and (
                pair["CONTCAR"]["front"] is not None or pair["CONTCAR"]["side"] is not None
            )

            if pos_ok and con_ok:
                self.structures.append(
                    {
                        "prefix": prefix,
                        "mode": "pair",
                        "poscar": pair["POSCAR"],
                        "contcar": pair["CONTCAR"],
                    }
                )
                continue

            single = data["single"]
            if single["top"] is not None and (single["front"] is not None or single["side"] is not None):
                self.structures.append({"prefix": prefix, "mode": "single", "single": single})
                continue

            print(f"警告: 结构 '{prefix}' 图像不完整，已跳过")

        if not self.structures:
            raise ValueError("未找到任何有效结构图像")

        print(f"已加载 {len(self.structures)} 个结构: {[s['prefix'] for s in self.structures]}")

    @staticmethod
    def _fit_image(img_path: str, x: float, y: float, w: float, h: float) -> Optional[Tuple[str, float, float, float, float]]:
        if not img_path or not os.path.exists(img_path):
            return None
        with Image.open(img_path) as im:
            iw, ih = im.size
        if iw <= 0 or ih <= 0 or w <= 0 or h <= 0:
            return None

        aspect = iw / ih
        target_w = w
        target_h = target_w / aspect
        if target_h > h:
            target_h = h
            target_w = target_h * aspect

        left = x + (w - target_w) / 2
        top = y + (h - target_h) / 2
        return img_path, left, top, target_w, target_h

    def _column_regions(self, n: int) -> List[Tuple[float, float, float, float]]:
        total_w = self.SLIDE_WIDTH - 2 * self.OUTER_PADDING - (n - 1) * self.COLUMN_GAP
        col_w = total_w / n
        regions = []
        for i in range(n):
            left = self.OUTER_PADDING + i * (col_w + self.COLUMN_GAP)
            regions.append((left, self.OUTER_PADDING, col_w, self.SLIDE_HEIGHT - 2 * self.OUTER_PADDING))
        return regions

    def _draw_pair_block(self, slide, struct: Dict, region: Tuple[float, float, float, float]):
        """画单个结构的左右对比块（POSCAR vs CONTCAR）。"""
        left, top, width, height = region

        title_h = height * self.TITLE_HEIGHT_RATIO
        title_x = left + width * self.TITLE_LEFT_OFFSET_RATIO
        title_y = top + height * self.TITLE_TOP_OFFSET_RATIO
        title_w = width * self.TITLE_WIDTH_RATIO
        title_box = slide.shapes.add_textbox(Inches(title_x), Inches(title_y), Inches(title_w), Inches(title_h))
        tf = title_box.text_frame
        tf.text = struct["prefix"]
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.size = Pt(self.TITLE_FONT_SIZE)
        tf.paragraphs[0].font.bold = True

        content_top = title_y + title_h + height * self.CONTENT_TOP_GAP_RATIO
        content_h = height - (content_top - top)

        panel_gap = width * self.PANEL_GAP_RATIO
        panel_w = (width - panel_gap) / 2
        left_panel_x = left
        right_panel_x = left + panel_w + panel_gap
        self._draw_one_panel(slide, left_panel_x, content_top, panel_w, content_h, "POSCAR", struct["poscar"])
        self._draw_one_panel(slide, right_panel_x, content_top, panel_w, content_h, "CONTCAR", struct["contcar"])

    def _draw_one_panel(self, slide, x: float, y: float, w: float, h: float, label: str, views: Dict[str, Optional[str]]):
        label_h = h * self.LABEL_HEIGHT_RATIO
        label_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(label_h))
        tf = label_box.text_frame
        tf.text = label
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.size = Pt(self.PANEL_LABEL_FONT_SIZE)
        tf.paragraphs[0].font.bold = True

        imgs_top = y + label_h + h * self.IMAGE_GAP_RATIO
        imgs_h = h - label_h - h * self.IMAGE_GAP_RATIO
        inner_w = w * self.IMAGE_WIDTH_RATIO
        inner_x = x + (w - inner_w) / 2

        gap_h = h * self.IMAGE_GAP_RATIO
        top_h = imgs_h * self.TOP_IMAGE_HEIGHT_RATIO
        bottom_h = imgs_h * self.BOTTOM_IMAGE_HEIGHT_RATIO
        total_h = top_h + bottom_h
        if total_h > 0:
            scale = imgs_h / total_h
            top_h *= scale
            bottom_h *= scale

        top_img = views.get("top")
        front_img_val = views.get("front")
        side_img_val = views.get("side")
        front_or_side = front_img_val or side_img_val

        p1 = self._fit_image(top_img, inner_x, imgs_top, inner_w, top_h - gap_h / 2)
        if p1:
            slide.shapes.add_picture(p1[0], Inches(p1[1]), Inches(p1[2]), width=Inches(p1[3]), height=Inches(p1[4]))

        bottom_y = imgs_top + top_h + gap_h
        if self.save_both_views and front_img_val and side_img_val:
            # 并排显示 front 和 side
            target_h = bottom_h - gap_h / 2
            with Image.open(front_img_val) as im_f, Image.open(side_img_val) as im_s:
                w_f, h_f = im_f.size
                w_s, h_s = im_s.size
            scaled_w_f = w_f * (target_h / h_f) if h_f > 0 else 0
            scaled_w_s = w_s * (target_h / h_s) if h_s > 0 else 0
            total_scaled_w = scaled_w_f + scaled_w_s
            if total_scaled_w > inner_w:
                scale = inner_w / total_scaled_w
                target_h *= scale
                scaled_w_f *= scale
                scaled_w_s *= scale
            start_x = inner_x + (inner_w - (scaled_w_f + scaled_w_s)) / 2
            p_front = self._fit_image(front_img_val, start_x, bottom_y, scaled_w_f, target_h)
            if p_front:
                slide.shapes.add_picture(p_front[0], Inches(p_front[1]), Inches(p_front[2]),
                                        width=Inches(p_front[3]), height=Inches(p_front[4]))
            p_side = self._fit_image(side_img_val, start_x + scaled_w_f, bottom_y, scaled_w_s, target_h)
            if p_side:
                slide.shapes.add_picture(p_side[0], Inches(p_side[1]), Inches(p_side[2]),
                                        width=Inches(p_side[3]), height=Inches(p_side[4]))
        else:
            p2 = self._fit_image(front_or_side, inner_x, bottom_y, inner_w, bottom_h - gap_h / 2)
            if p2:
                slide.shapes.add_picture(p2[0], Inches(p2[1]), Inches(p2[2]), width=Inches(p2[3]), height=Inches(p2[4]))

    def _draw_single_block(self, slide, struct: Dict, region: Tuple[float, float, float, float]):
        """单结构模式：显示 top + (可能并排的 front/side)。"""
        left, top, width, height = region

        title_h = height * self.TITLE_HEIGHT_RATIO
        title_x = left + width * self.TITLE_LEFT_OFFSET_RATIO
        title_y = top + height * self.TITLE_TOP_OFFSET_RATIO
        title_w = width * self.TITLE_WIDTH_RATIO

        title_box = slide.shapes.add_textbox(Inches(title_x), Inches(title_y), Inches(title_w), Inches(title_h))
        tf = title_box.text_frame
        tf.text = struct["prefix"]
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].font.size = Pt(self.TITLE_FONT_SIZE)
        tf.paragraphs[0].font.bold = True

        views = struct["single"]
        content_top = title_y + title_h + height * self.CONTENT_TOP_GAP_RATIO
        content_h = height - (content_top - top)
        gap = content_h * self.IMAGE_GAP_RATIO
        inner_w = width * self.IMAGE_WIDTH_RATIO
        inner_x = left + (width - inner_w) / 2

        top_img = views.get("top")
        front_img = views.get("front")
        side_img = views.get("side")

        # 上方 top 视图
        top_h = content_h * self.TOP_IMAGE_HEIGHT_RATIO
        p_top = self._fit_image(top_img, inner_x, content_top, inner_w, top_h - gap / 2)
        if p_top:
            slide.shapes.add_picture(p_top[0], Inches(p_top[1]), Inches(p_top[2]),
                                     width=Inches(p_top[3]), height=Inches(p_top[4]))

        bottom_y = content_top + top_h + gap
        bottom_h = content_h - top_h - gap

        if self.save_both_views and front_img and side_img:
            # 水平并排 front 和 side
            target_h = bottom_h
            with Image.open(front_img) as im_f, Image.open(side_img) as im_s:
                w_f, h_f = im_f.size
                w_s, h_s = im_s.size
            scaled_w_f = w_f * (target_h / h_f) if h_f > 0 else 0
            scaled_w_s = w_s * (target_h / h_s) if h_s > 0 else 0
            total_scaled_w = scaled_w_f + scaled_w_s
            if total_scaled_w > inner_w:
                scale = inner_w / total_scaled_w
                target_h *= scale
                scaled_w_f *= scale
                scaled_w_s *= scale
            start_x = inner_x + (inner_w - (scaled_w_f + scaled_w_s)) / 2
            p_front = self._fit_image(front_img, start_x, bottom_y, scaled_w_f, target_h)
            if p_front:
                slide.shapes.add_picture(p_front[0], Inches(p_front[1]), Inches(p_front[2]),
                                         width=Inches(p_front[3]), height=Inches(p_front[4]))
            p_side = self._fit_image(side_img, start_x + scaled_w_f, bottom_y, scaled_w_s, target_h)
            if p_side:
                slide.shapes.add_picture(p_side[0], Inches(p_side[1]), Inches(p_side[2]),
                                         width=Inches(p_side[3]), height=Inches(p_side[4]))
        else:
            lower_img = front_img or side_img
            p_lower = self._fit_image(lower_img, inner_x, bottom_y, inner_w, bottom_h - gap / 2)
            if p_lower:
                slide.shapes.add_picture(p_lower[0], Inches(p_lower[1]), Inches(p_lower[2]),
                                         width=Inches(p_lower[3]), height=Inches(p_lower[4]))

    def _draw_column_dividers(self, slide, regions: List[Tuple[float, float, float, float]]):
        """在多结构布局中绘制列间竖向分割线。"""
        if not self.SHOW_COLUMN_DIVIDER:
            return
        if len(regions) <= 1:
            return

        for i in range(len(regions) - 1):
            left, top, width, height = regions[i]
            next_left, _, _, _ = regions[i + 1]
            gap = next_left - (left + width)
            x = left + width + gap / 2 + self.DIVIDER_X_OFFSET
            y = top + height * self.DIVIDER_Y_OFFSET_RATIO
            h = height * self.DIVIDER_HEIGHT_RATIO

            line = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x),
                Inches(y),
                Inches(0),
                Inches(h),
            )
            line.line.color.rgb = RGBColor(*self.DIVIDER_COLOR_RGB)
            line.line.width = Pt(self.DIVIDER_WIDTH_PT)

    def generate_ppt(self, output_path: str, structures_per_slide: int):
        if structures_per_slide not in (1, 2, 3):
            raise ValueError("structures_per_slide 必须是 1、2 或 3")

        out_dir = os.path.dirname(output_path)
        if out_dir and not os.path.exists(out_dir):
            os.makedirs(out_dir)

        prs = Presentation()
        prs.slide_width = Inches(self.SLIDE_WIDTH)
        prs.slide_height = Inches(self.SLIDE_HEIGHT)
        blank = prs.slide_layouts[6]

        for i in range(0, len(self.structures), structures_per_slide):
            chunk = self.structures[i : i + structures_per_slide]
            slide = prs.slides.add_slide(blank)
            regions = self._column_regions(len(chunk))
            self._draw_column_dividers(slide, regions)

            for struct, region in zip(chunk, regions):
                if struct["mode"] == "pair":
                    self._draw_pair_block(slide, struct, region)
                else:
                    self._draw_single_block(slide, struct, region)

        prs.save(output_path)
        print(f"PPT 已生成: {output_path}")